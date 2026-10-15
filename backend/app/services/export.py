from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.reaction_route import ExtractionResult, PipelineStageStatus


def _entity_names(entities: list[Any]) -> str:
    parts: list[str] = []
    for ent in entities or []:
        name = getattr(ent, "name", None) or (ent.get("name") if isinstance(ent, dict) else "")
        amount = getattr(ent, "amount", None) or (ent.get("amount") if isinstance(ent, dict) else "")
        smiles = getattr(ent, "smiles", None) or (ent.get("smiles") if isinstance(ent, dict) else "")
        if not name:
            continue
        label = str(name)
        if amount and amount not in ("N/A", ""):
            label = f"{label} ({amount})"
        if smiles and smiles != name:
            label = f"{label} [{smiles}]"
        parts.append(label)
    return "; ".join(parts) if parts else ""


def collect_warnings(result: ExtractionResult) -> list[str]:
    warnings: list[str] = []
    if result.error:
        warnings.append(result.error)
    route = result.route
    if route and route.warnings:
        warnings.extend([w for w in route.warnings if w])
    for stage in result.stages:
        if stage.status == PipelineStageStatus.error and stage.detail:
            warnings.append(f"{stage.label}: {stage.detail}")
        elif stage.id == "cache" and stage.detail:
            warnings.append(stage.detail)
    seen: set[str] = set()
    unique: list[str] = []
    for w in warnings:
        key = w.strip()
        if key and key not in seen:
            seen.add(key)
            unique.append(key)
    return unique


def write_result_json(result: ExtractionResult, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(result.model_dump_json(indent=2), encoding="utf-8")
    return path


def write_batch_json(results: list[ExtractionResult], batch_id: str, path: Path) -> Path:
    payload = {
        "batch_id": batch_id,
        "file_count": len(results),
        "results": [json.loads(r.model_dump_json()) for r in results],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def _append_steps_sheet(ws, results: list[ExtractionResult]) -> None:
    headers = [
        "filename",
        "job_id",
        "route_title",
        "target_name",
        "target_class",
        "step_index",
        "step_label",
        "reactants",
        "reagents",
        "catalysts_ligands",
        "solvents",
        "products",
        "intermediates",
        "temperature",
        "time",
        "atmosphere",
        "other_conditions",
        "yield",
        "ee_or_dr",
        "source",
        "confidence",
        "provenance_pages",
        "notes",
        "evidence",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for result in results:
        route = result.route
        if not route or not route.steps:
            ws.append(
                [
                    result.filename,
                    result.job_id,
                    route.title if route else "",
                    route.target_name if route else "",
                    route.target_class if route else "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    result.error or "No steps extracted",
                    "",
                ]
            )
            continue
        for step in route.steps:
            cond = step.conditions
            ws.append(
                [
                    result.filename,
                    result.job_id,
                    route.title,
                    route.target_name,
                    route.target_class,
                    step.step_index,
                    step.step_label,
                    _entity_names(step.reactants),
                    _entity_names(step.reagents),
                    _entity_names(step.catalysts + step.ligands),
                    _entity_names(step.solvents),
                    _entity_names(step.products),
                    _entity_names(step.intermediates),
                    cond.temperature if cond else "N/A",
                    cond.time if cond else "N/A",
                    cond.atmosphere if cond else "N/A",
                    cond.other if cond else "N/A",
                    step.yield_value,
                    step.ee_or_dr,
                    step.source or "",
                    step.confidence if step.confidence is not None else "",
                    ", ".join(str(r.page) for r in step.evidence_refs if r.page),
                    step.notes,
                    step.evidence[:800] if step.evidence else "",
                ]
            )


def _append_routes_sheet(ws, results: list[ExtractionResult]) -> None:
    ws.append(
        [
            "filename",
            "job_id",
            "title",
            "target_name",
            "target_class",
            "document_type",
            "n_steps",
            "from_cache",
            "warnings",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        ws.append(
            [
                result.filename,
                result.job_id,
                route.title if route else "",
                route.target_name if route else "",
                route.target_class if route else "",
                route.document_type if route else "",
                len(route.steps) if route else 0,
                "yes" if result.from_cache else "no",
                " | ".join(collect_warnings(result)),
            ]
        )


def _append_compounds_sheet(ws, results: list[ExtractionResult]) -> None:
    ws.append(
        [
            "filename",
            "job_id",
            "step_index",
            "step_label",
            "role",
            "name",
            "amount",
            "smiles",
            "cas",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        if not route:
            continue
        for step in route.steps:
            for bucket in (
                step.reactants,
                step.reagents,
                step.catalysts,
                step.ligands,
                step.solvents,
                step.intermediates,
                step.products,
            ):
                for ent in bucket:
                    ws.append(
                        [
                            result.filename,
                            result.job_id,
                            step.step_index,
                            step.step_label,
                            ent.role,
                            ent.name,
                            ent.amount or "",
                            ent.smiles or "",
                            ent.cas or "",
                        ]
                    )


def _append_route_graph_sheet(ws, results: list[ExtractionResult]) -> None:
    ws.append(
        [
            "filename",
            "job_id",
            "node_id",
            "name",
            "canonical_name",
            "compound_ids",
            "step_indices",
            "role",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        if not route or not route.route_graph:
            continue
        for node in route.route_graph.nodes:
            ws.append(
                [
                    result.filename,
                    result.job_id,
                    node.node_id,
                    node.name,
                    node.canonical_name,
                    ", ".join(node.compound_ids),
                    ", ".join(str(i) for i in node.step_indices),
                    node.role,
                ]
            )


def _append_route_edges_sheet(ws, results: list[ExtractionResult]) -> None:
    ws.append(
        [
            "filename",
            "job_id",
            "from_node",
            "to_node",
            "step_index",
            "relation",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        if not route or not route.route_graph:
            continue
        for edge in route.route_graph.edges:
            ws.append(
                [
                    result.filename,
                    result.job_id,
                    edge.from_node,
                    edge.to_node,
                    edge.step_index,
                    edge.relation,
                ]
            )


def _append_compound_links_sheet(ws, results: list[ExtractionResult]) -> None:
    ws.append(
        [
            "filename",
            "job_id",
            "compound_id",
            "name",
            "role",
            "example_label",
            "step_indices",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        if not route:
            continue
        for link in route.compound_links or []:
            ws.append(
                [
                    result.filename,
                    result.job_id,
                    link.compound_id,
                    link.name,
                    link.role,
                    link.example_label,
                    ", ".join(str(i) for i in link.step_indices),
                ]
            )


def build_excel_workbook(results: list[ExtractionResult]) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(
        [
            "filename",
            "job_id",
            "status",
            "from_cache",
            "pages",
            "figures",
            "steps",
            "target_name",
            "target_class",
            "title",
            "error",
            "warnings",
        ]
    )
    for cell in summary[1]:
        cell.font = Font(bold=True)

    for result in results:
        doc = result.document
        route = result.route
        failed = result.error or any(s.status == PipelineStageStatus.error for s in result.stages)
        warnings = collect_warnings(result)
        summary.append(
            [
                result.filename,
                result.job_id,
                "error" if failed else "ok",
                "yes" if result.from_cache else "no",
                doc.page_count if doc else "",
                doc.figure_count if doc else "",
                len(route.steps) if route else 0,
                route.target_name if route else "",
                route.target_class if route else "",
                route.title if route else "",
                result.error or "",
                " | ".join(warnings),
            ]
        )

    routes_ws = wb.create_sheet("Routes")
    _append_routes_sheet(routes_ws, results)

    steps_ws = wb.create_sheet("Steps")
    _append_steps_sheet(steps_ws, results)

    compounds_ws = wb.create_sheet("Compounds")
    _append_compounds_sheet(compounds_ws, results)

    graph_ws = wb.create_sheet("RouteGraph")
    _append_route_graph_sheet(graph_ws, results)

    edges_ws = wb.create_sheet("RouteEdges")
    _append_route_edges_sheet(edges_ws, results)

    links_ws = wb.create_sheet("CompoundLinks")
    _append_compound_links_sheet(links_ws, results)

    warn_ws = wb.create_sheet("Warnings")
    warn_ws.append(["filename", "job_id", "source", "warning"])
    for cell in warn_ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        route = result.route
        if result.error:
            warn_ws.append([result.filename, result.job_id, "error", result.error])
        if route:
            for w in route.warnings or []:
                if w:
                    warn_ws.append([result.filename, result.job_id, "route", w])
        for stage in result.stages:
            if stage.status == PipelineStageStatus.error and stage.detail:
                warn_ws.append([result.filename, result.job_id, stage.id, stage.detail])
            elif stage.id == "cache" and stage.detail:
                warn_ws.append([result.filename, result.job_id, "cache", stage.detail])
        if not result.error and not (route and route.warnings) and not any(
            s.status == PipelineStageStatus.error for s in result.stages
        ):
            if not route or not route.steps:
                warn_ws.append(
                    [result.filename, result.job_id, "extract", "No reaction steps extracted"]
                )

    return wb


def write_excel(results: list[ExtractionResult], path: Path) -> Path:
    wb = build_excel_workbook(results)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def excel_bytes(results: list[ExtractionResult]) -> bytes:
    buf = BytesIO()
    build_excel_workbook(results).save(buf)
    return buf.getvalue()
