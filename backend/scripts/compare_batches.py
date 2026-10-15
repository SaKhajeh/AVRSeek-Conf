"""Compare two AVRSeek batch Excel exports."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


def sheet_to_rows(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def main() -> None:
    new_path = Path(r"c:\Users\galin\Downloads\avrseek_batch_f4af663d45de.xlsx")
    old_path = Path(
        r"C:\Users\galin\Downloads\DeepSynthesis-ReactionSeek-6d42c4a\AVRSeek\backend\outputs\batches\0e7175dfc9d9\batch.xlsx"
    )
    for p in (new_path, old_path):
        if not p.exists():
            print(f"MISSING: {p}")
            return

    new_wb = load_workbook(new_path, read_only=True, data_only=True)
    old_wb = load_workbook(old_path, read_only=True, data_only=True)
    print("NEW sheets:", new_wb.sheetnames)
    print("OLD sheets:", old_wb.sheetnames)
    print()

    for sheet in ("Summary", "Steps", "RouteGraph", "RouteEdges", "CompoundLinks", "Warnings"):
        new_rows = sheet_to_rows(new_path, sheet)
        old_rows = sheet_to_rows(old_path, sheet)
        print(f"--- {sheet} ---")
        print(f"  old rows: {max(0, len(old_rows)-1)}, new rows: {max(0, len(new_rows)-1)}")
        if new_rows:
            print(f"  new headers: {new_rows[0]}")
        if sheet == "Steps" and len(new_rows) > 1:
            hdr = new_rows[0]
            for row in new_rows[1:]:
                d = dict(zip(hdr, row))
                print(
                    f"  step {d.get('step_index')}: {str(d.get('step_label',''))[:70]} | "
                    f"conf={d.get('confidence')} | src={d.get('source')}"
                )
        if sheet == "Summary" and len(new_rows) > 1:
            print(f"  new summary: {new_rows[1]}")
        if sheet == "Warnings" and len(new_rows) > 1:
            for row in new_rows[1:]:
                print(f"  warn: {row}")
        print()

    # JSON batch if present
    batch_dir = Path(r"C:\Users\galin\Downloads\DeepSynthesis-ReactionSeek-6d42c4a\AVRSeek\backend\outputs\batches")
    for bid in ("f4af663d45de", "0e7175dfc9d9"):
        jpath = batch_dir / bid / "batch.json"
        if jpath.exists():
            data = json.loads(jpath.read_text(encoding="utf-8"))
            r = data["results"][0] if data.get("results") else {}
            route = r.get("route") or {}
            stages = {s["id"]: s for s in r.get("stages", [])}
            debug = r.get("extraction_debug") or {}
            print(f"=== JSON {bid} ===")
            print(f"  steps: {len(route.get('steps', []))}")
            print(f"  route_graph nodes: {len((route.get('route_graph') or {}).get('nodes', []))}")
            print(f"  compound_links: {len(route.get('compound_links', []))}")
            for sid in ("image_scheme", "table_extract", "reconcile", "standardize"):
                st = stages.get(sid, {})
                print(f"  stage {sid}: {st.get('status')} — {str(st.get('detail',''))[:100]}")
            print(f"  figure_pages_vlm: {debug.get('figure_pages_vlm')}")
            print(f"  route_graph_nodes debug: {debug.get('route_graph_nodes')}")
            print()


    for bid in ("0e7175dfc9d9", "f4af663d45de"):
        r, route = load(bid)
        print("===", bid, "===")
        print("job:", r["job_id"])
        for s in route["steps"]:
            smiles = sum(
                1
                for b in (
                    "reactants",
                    "reagents",
                    "catalysts",
                    "ligands",
                    "solvents",
                    "products",
                    "intermediates",
                )
                for e in s.get(b, [])
                if e.get("smiles")
            )
            label = (s.get("step_label") or "")[:75]
            print(
                f"  {s['step_index']}: {label} | conf={s.get('confidence')} "
                f"src={s.get('source')} smiles={smiles} yield={s.get('yield_value')}"
            )
        rg = route.get("route_graph") or {}
        print("  graph:", len(rg.get("nodes", [])), "nodes,", len(rg.get("edges", [])), "edges")
        print("  warnings:", len(route.get("warnings", [])))
        print()


if __name__ == "__main__":
    main()
